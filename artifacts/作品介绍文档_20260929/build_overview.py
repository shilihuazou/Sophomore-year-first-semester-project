from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(r"D:\CangQiongPj\artifacts\作品介绍文档_20260929\吉隆口岸泥石流监测救援部署系统系统概述.docx")
SONG = "宋体"
HEI = "黑体"


def set_font(run, name=SONG, size=10.5, bold=False, color="000000"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    fonts = r_pr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, fonts)
    for attribute in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{attribute}"), name)


def paragraph(doc, text, first_indent=True, after=7):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(21) if first_indent else Pt(0)
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(after)
    set_font(p.add_run(text))
    return p


def heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(5)
    set_font(p.add_run(text), HEI, 12, True)


def numbered(doc, index, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.65)
    p.paragraph_format.first_line_indent = Cm(-0.65)
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(6)
    set_font(p.add_run(f"{index}  {label}。"), SONG, 10.5, True)
    set_font(p.add_run(text))


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

    normal = doc.styles["Normal"]
    normal.font.name = SONG
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), SONG)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(12)
    title.paragraph_format.space_after = Pt(8)
    set_font(title.add_run("吉隆口岸泥石流监测救援部署系统"), HEI, 18, True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(20)
    set_font(subtitle.add_run("系统概述"), SONG, 12)

    heading(doc, "一  背景介绍")
    paragraph(
        doc,
        "本作品以吉隆口岸泥石流灾害应对为场景，围绕空间资料查看、现场点位标记和救援任务协同展开。项目利用赛事提供的遥感影像及吉隆县冰川、冰湖、冻土、历史灾害区等原始资料，基于 KQGIS Desktop 与 KQGIS Server 进行地图服务发布和二次开发，使管理人员能够在同一地图界面查看环境信息、部署队伍并跟踪处置进度。",
    )

    heading(doc, "二  主要功能")
    numbered(
        doc,
        "1",
        "空间数据展示",
        "智慧大屏以吉隆口岸影像为核心地图，可切换 9 景多时相影像；左侧展示冰川、冰湖、冻土和历史灾害区的原始数据卡片及图表。点击卡片即可切换至对应的吉隆县专题地图。",
    )
    numbered(
        doc,
        "2",
        "救援点位调度",
        "管理员在口岸地图上点击创建点位，在点位旁选择具体救援队伍；分配时自动生成唯一任务编号。后台可维护抢险救援、医疗治疗、电力保障、地质监测和群众安置五类队伍的组长及成员信息，大屏同步显示各类队伍的可用数量。",
    )
    numbered(
        doc,
        "3",
        "现场任务上报",
        "组长使用登记的手机号登录移动端，仅查看本组任务，可上传现场图片并推进任务阶段。任务进度按已分配 0%、任务中 50%、任务完成 100% 展示；大屏同步回显阶段、图片和点位图标状态。",
    )

    heading(doc, "三  作品特点")
    paragraph(
        doc,
        "作品把地图展示与救援业务连接为“管理员落点—分配队伍—组长上报—大屏回显”的闭环。未分配点位使用水滴图标，分配后显示对应组别的圆形矢量图标，完成后变为绿底，便于快速辨认。专题卡片的数字和图表直接取自赛事原始文件已有记录与字段，卡片标明数据来源，且与已发布图层对应；业务点位和上报内容则单独保存，避免改动原始空间资料。系统由 Vue 3、OpenLayers、Django 5 和 KQGIS Server 协同实现。",
        after=0,
    )

    doc.core_properties.title = "吉隆口岸泥石流监测救援部署系统系统概述"
    doc.core_properties.subject = "苍穹杯现场答辩阅览材料"
    doc.core_properties.author = "项目团队"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
