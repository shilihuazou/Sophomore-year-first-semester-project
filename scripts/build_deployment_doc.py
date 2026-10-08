"""Generate the reviewed Windows deployment guide for the competition project."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "artifacts" / "部署说明文档_20260929" / "吉隆口岸泥石流监测救援部署系统部署说明.docx"


def set_font(style, name="宋体", size=11, bold=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), name)


def set_cell_text(cell, text, bold=False, color="000000", size=9.5):
    cell.text = ""
    para = cell.paragraphs[0]
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing = 1.12
    run = para.add_run(str(text))
    run.bold = bold
    run.font.name = "宋体"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    fonts = run._element.get_or_add_rPr().rFonts
    if fonts is not None:
        fonts.set(qn("w:eastAsia"), "宋体")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def border_cell(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D9D9D9")
    mar = tcpr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for edge, value in (("top", 75), ("bottom", 75), ("left", 95), ("right", 95)):
        tag = qn(f"w:{edge}")
        item = mar.find(tag)
        if item is None:
            item = OxmlElement(f"w:{edge}")
            mar.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, width in enumerate(widths):
        t.columns[i].width = Inches(width)
    for i, heading in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.width = Inches(widths[i])
        set_cell_text(cell, heading, bold=True, color="FFFFFF", size=9.5)
        shade(cell, "25435B")
        border_cell(cell)
    trpr = t.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    trpr.append(repeat)
    no_split = OxmlElement("w:cantSplit")
    trpr.append(no_split)
    for row_i, values in enumerate(rows):
        row = t.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        cells = row.cells
        for i, value in enumerate(values):
            cells[i].width = Inches(widths[i])
            set_cell_text(cells[i], value)
            if row_i % 2:
                shade(cells[i], "F4F7F9")
            border_cell(cells[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def para(doc, text="", style=None):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


def lead(doc, label, text):
    p = doc.add_paragraph()
    p.add_run(label).bold = True
    p.add_run(text)
    return p


def bullets(doc, items):
    for item in items:
        para(doc, item, "List Bullet")


def numbered(doc, items):
    for index, item in enumerate(items, 1):
        p = para(doc, f"{index}.  {item}")
        p.paragraph_format.left_indent = Inches(0.23)
        p.paragraph_format.first_line_indent = Inches(-0.23)
        p.paragraph_format.space_after = Pt(3)


def code(doc, lines):
    for line in lines:
        p = para(doc, line if line else " ", "Code")
        p.paragraph_format.keep_with_next = line != lines[-1]


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.68)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
set_font(styles["Normal"])
styles["Normal"].paragraph_format.line_spacing = 1.16
styles["Normal"].paragraph_format.space_after = Pt(5)
set_font(styles["Title"], "黑体", 20, True)
styles["Title"].paragraph_format.space_after = Pt(12)
set_font(styles["Heading 1"], "黑体", 14, True)
styles["Heading 1"].paragraph_format.space_before = Pt(13)
styles["Heading 1"].paragraph_format.space_after = Pt(6)
styles["Heading 1"].paragraph_format.keep_with_next = True
set_font(styles["Heading 2"], "黑体", 11.5, True)
styles["Heading 2"].paragraph_format.space_before = Pt(9)
styles["Heading 2"].paragraph_format.space_after = Pt(4)
styles["Heading 2"].paragraph_format.keep_with_next = True
set_font(styles["List Bullet"], "宋体", 10.5)
styles["List Bullet"].paragraph_format.left_indent = Inches(0.23)
styles["List Bullet"].paragraph_format.first_line_indent = Inches(-0.16)
styles["List Bullet"].paragraph_format.space_after = Pt(3)
set_font(styles["List Number"], "宋体", 10.5)
styles["List Number"].paragraph_format.left_indent = Inches(0.23)
styles["List Number"].paragraph_format.first_line_indent = Inches(-0.16)
styles["List Number"].paragraph_format.space_after = Pt(3)
if "Code" not in styles:
    styles.add_style("Code", 1)
set_font(styles["Code"], "Consolas", 9)
styles["Code"].paragraph_format.left_indent = Inches(0.24)
styles["Code"].paragraph_format.line_spacing = 1.0
styles["Code"].paragraph_format.space_after = Pt(2)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("吉隆口岸泥石流监测救援部署系统  ·  部署说明  ·  2026 年 9 月")
run.font.name = "宋体"
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(110, 110, 110)

title = para(doc, "吉隆口岸泥石流监测救援部署系统部署说明", "Title")
title_ppr = title._p.get_or_add_pPr()
title_borders = OxmlElement("w:pBdr")
title_bottom = OxmlElement("w:bottom")
title_bottom.set(qn("w:val"), "nil")
title_borders.append(title_bottom)
title_ppr.append(title_borders)
para(doc, "适用对象：比赛演示环境的服务器部署人员    文档版本：2026 年 9 月 29 日")
lead(doc, "先读结论。", "本工程可在 Windows x64 主机按原有目录结构迁移并通过一键脚本启动管理员大屏、组长任务端、Django API 与 KQGIS Server。当前启动器使用 Vite 开发服务器和 Django runserver，适合受控的比赛演示；若要长期公开运行，须先完成正式 Web 服务部署与账户安全加固。")

doc.add_heading("1 系统组成与部署边界", 1)
para(doc, "浏览器访问两个独立的 Vue 页面；前端开发服务器把 /api 请求转发至同机 Django，管理员端还把 /kqgis 请求转发至同机 KQGIS。Django 使用 SQLite 保存账号、队伍、点位与上报，现场照片存放在 media 目录；KQGIS 地图服务读取比赛方原始 GIS 资料。")
table(doc, ["组件", "本机端口与作用", "对外暴露原则"], [
    ("管理员端", "5173；登录、智慧大屏、调度后台", "局域网演示可开放；公网宜仅经 HTTPS 代理访问"),
    ("组长端", "5174；任务查看、阶段上报、照片上传", "与管理员端相同；手机从这里访问"),
    ("Django", "127.0.0.1:8000；/api 接口", "不直接开放公网，经过前端代理"),
    ("KQGIS Server", "9600；地图元数据、WMS", "供应商程序的监听范围须实测，并用防火墙限制外部访问"),
], [1.12, 3.18, 2.75])
lead(doc, "核心运行数据。", "应把 app/backend/db.sqlite3 与 app/backend/media 作为同一批次迁移和备份；把 Gis/kqmsdata、kqgis_publish 及“尼泊尔泥石流巨灾”原始数据随工程保留。只复制前端 dist 或源码，无法恢复已发布地图和任务记录。")

doc.add_heading("2 部署前检查", 1)
table(doc, ["项目", "要求与检查点"], [
    ("操作系统", "Windows x64；当前 KQGIS Server 为 Windows 发行包。部署账号需能读写工程与日志目录。"),
    ("Python", "Python 3.10 或更新，安装 py 启动器；requirements.txt 固定 Django 5.2.17。不要把 KQGIS 自带 Python 当作应用 Python。"),
    ("Node.js", "Vite 8.3.0 要求 Node ^20.19.0 或 >=22.12.0；同时需要 npm。"),
    ("KQGIS", "保留 Gis 内程序、kqmsdata 和有效授权；部署前检查授权与服务状态。"),
    ("目录与磁盘", "建议完整工程仍放 D:\\CangQiongPj；地图工程可能保存绝对数据路径。按原始数据实际体量预留空间。"),
    ("端口", "5173、5174、8000、9600 不应被无关程序占用；公网通常仅放通 80/443。"),
], [1.35, 5.7])
p = para(doc, "在新主机 PowerShell 中先检查运行时和端口；部署包尚未复制时，只执行前两条即可：")
p.paragraph_format.keep_with_next = True
code(doc, [
    "py --version",
    "node --version; npm --version",
    "Get-NetTCPConnection -State Listen -LocalPort 5173,5174,8000,9600 -ErrorAction SilentlyContinue",
])
para(doc, "若必须改变项目根路径，应先在 KQGIS Desktop/Server 中修正原始图层连接并重新发布、验收地图服务，不要仅改前端 URL。授权文件及地图源数据不得随意改名。")

doc.add_heading("3 复制工程与保留数据", 1)
numbered(doc, [
    "在原主机记录管理员端、组长端、KQGIS 服务和数据库的当前状态；维护窗口内停止写入，避免复制 SQLite 时仍有任务上报。",
    "将 D:\\CangQiongPj 完整复制到新主机同名路径。至少保留 app、Gis、kqgis_publish、尼泊尔泥石流巨灾、licenses、scripts 和根目录启动脚本。",
    "核对 app/backend/db.sqlite3 与 app/backend/media/marker_reports 是否同时存在；已有账号、队伍、点位和照片依赖这两处内容。",
    "核对 Gis/kqmsdata 与 KQGIS 可执行程序、授权文件；缺少或路径失效时，页面可打开但地图可能为空。",
    "node_modules 和 dist 可重新生成；若目标主机无外网安装能力，应先准备可用的离线依赖，不要依赖启动器自动下载。",
])
lead(doc, "注意。", "一键启动会执行 Django 数据库迁移。迁移现有 db.sqlite3 之前，先保留一份独立备份；不要在已有比赛数据的库上运行重置脚本或以空库覆盖。")

doc.add_heading("4 配置访问地址与安全参数", 1)
doc.add_heading("4.1 本机与局域网演示", 2)
para(doc, "app/network-config.ps1 默认 BindAddress 为 0.0.0.0，两个前端可被同网段设备访问。启动器自动识别本机 IPv4，并将其加入 Vite 允许主机及 Django CSRF 配置；手机与服务器处于同一可互访网络时，可使用启动输出的 LAN 地址。Windows 防火墙只放行演示所需的 5173/5174，不要为 8000/9600 建立公网入站放行。")
doc.add_heading("4.2 域名和 HTTPS 演示", 2)
para(doc, "先在 DNS、云安全组或路由器中准备可达的 HTTPS 入口，再在 app/network-config.ps1 的 PublicOrigins 填入浏览器实际访问的来源。来源只含协议、域名和可选端口，不包含路径、斜杠后的页面或 #/login。示例如下，域名必须替换为实际值：")
code(doc, [
    "PublicOrigins = @('https://admin.example.com', 'https://team.example.com')",
])
p = para(doc, "同机 HTTPS 反向代理可将管理员域名转发至 127.0.0.1:5173、组长域名转发至 127.0.0.1:5174，并支持 WebSocket。此模式可把 APP_BIND_ADDRESS 设为 127.0.0.1；若代理在另一台主机，则保持可达的监听地址并用防火墙只允许代理主机访问 5173/5174。浏览器的 /api 和管理员端的 /kqgis 仍由 Vite 代理到本机 8000/9600。")
p.paragraph_format.keep_together = True
p = para(doc, "短期 HTTPS 演示可在启动脚本所在的同一个 PowerShell 会话中设置环境变量，再启动服务；请用随机密钥替换占位符，不要把真实密钥写入交付文档或提交到仓库：")
p.paragraph_format.keep_with_next = True
code(doc, [
    "$env:DJANGO_SECRET_KEY = '<足够长的随机密钥>'",
    "$env:DJANGO_DEBUG = '0'",
    "$env:DJANGO_COOKIE_SECURE = '1'",
    "$env:DJANGO_USE_X_FORWARDED_HOST = '1'",
    "$env:DJANGO_TRUST_PROXY_PROTO = '1'",
    "& 'D:\\CangQiongPj\\scripts\\start.ps1'",
])
para(doc, "仅在可信代理确实设置并清理 X-Forwarded-Proto 等转发头时启用相应代理选项。纯 HTTP 局域网演示不要启用 DJANGO_COOKIE_SECURE=1，否则登录 Cookie 可能无法正常工作。修改来源或安全环境变量后，重新运行启动器使各服务使用新配置。")
lead(doc, "公网发布前的阻断项。", "现有管理员演示口令及组长默认口令 123 不适合互联网公开；组长端当前没有完整的逐人改密和生产级权限管理流程。未完成唯一密码、访问限制与正式服务化改造前，应将入口限定在评委网络、VPN 或受控临时演示环境。")

doc.add_heading("5 一键启动与首次初始化", 1)
numbered(doc, [
    "在工程根目录双击“一键启动全部环境.cmd”，或从 PowerShell 运行 powershell -ExecutionPolicy Bypass -File .\\scripts\\start.ps1。",
    "脚本检查 py、node、npm；缺少 Django 或 Vite 依赖时尝试安装，随后执行 manage.py migrate。首次执行需要网络或已备齐的离线依赖。",
    "启动器以隐藏窗口运行 8000、5173、5174，并通过 KQMS_DATA_PATH 指向本工程 Gis 启动 KQGIS；最长约 180 秒检测四项服务。",
    "看到 frontend、teamFrontend、backend、kqgis 均 ready 后，再打开页面。命令窗口关闭后子进程仍会运行；重启服务器后需要重新启动。",
])
code(doc, [
    "管理员端  http://127.0.0.1:5173/#/login",
    "组长端    http://127.0.0.1:5174/",
])
para(doc, "已有 db.sqlite3 时沿用原账号、队伍与点位，不运行 seed_demo。只有刻意创建全新空库时，才在 app/backend 执行 py manage.py seed_demo；该命令首次创建 admin 并打印随机密码，重复执行不会重置已有密码。新空库不会自动恢复六支演示队伍和历史任务。")
para(doc, "启动器会重启它能识别的 8000/5173/5174 旧进程；若端口被其他程序占用会停止并报错。9600 已有监听时可能沿用现有服务，须确认它确为本项目 KQGIS，不能只凭端口开放判断地图正确。")

doc.add_heading("6 部署验收", 1)
doc.add_heading("6.1 接口和端口", 2)
p = para(doc, "从服务器本机执行以下只读检查。前四项应返回可解析内容；如果 KQGIS 元数据返回的是 HTML 启动页或错误信息，应继续检查服务状态和授权。")
p.paragraph_format.keep_with_next = True
code(doc, [
    "Invoke-WebRequest http://127.0.0.1:5173/ -UseBasicParsing",
    "Invoke-WebRequest http://127.0.0.1:5174/ -UseBasicParsing",
    "Invoke-WebRequest http://127.0.0.1:8000/api/csrf/ -UseBasicParsing",
    "$gis = 'http://127.0.0.1:9600/kqgis/rest/services'",
    "Invoke-RestMethod \"$gis/jilong_port_multitemporal/map/getmetadata\"",
    "Invoke-RestMethod \"$gis/jilong_county_atlas/map/getmetadata\"",
])
doc.add_heading("6.2 浏览器端到端检查", 2)
numbered(doc, [
    "管理员登录后打开智慧大屏，确认吉隆口岸实景影像、左侧真实数据卡和地图上方队伍数量均显示。",
    "切换多时相影像；点击冰川、冰湖、冻土或历史灾害卡片，确认吉隆县专题图出图，再返回口岸。",
    "检查现有点位的分配弹窗和已分配图标；不要在正式数据上为验收随意新增或删除任务。",
    "进入调度后台，确认队伍资料和组长联系方式与迁移前一致。",
    "用受控测试账号在组长端核对任务编号、阶段与照片入口；如需实际上传测试，应先约定测试点位与回滚方式。",
    "从目标手机网络访问最终域名或 LAN 地址，确认登录、图片显示、接口请求和页面刷新均正常。",
])
para(doc, "本说明部署的组长端是 app/team-frontend 中的 Web 页面。若另有独立开发的手机应用安装包，还需单独配置它的接口域名，并针对会话、CSRF 和照片上传做真机联调；本工程的一键脚本不会自动发布该安装包。")
para(doc, "上线验收应另从外部网络确认仅预定入口可达，并验证 8000/9600 不可从外部直接访问。KQGIS 厂商服务可能监听多个网卡，不能仅依据启动器注释推断其暴露范围。")

doc.add_heading("7 日常运行、备份与更新", 1)
lead(doc, "日志与重启。", "应用启动日志在 app/.runtime/，按 backend、frontend、team-frontend、kqgis 和时间戳区分 stdout/stderr；KQGIS 厂商日志还在 Gis/kqmsdata/logs 与发行包的 kqwebservice/logs。系统重启后重新运行启动器。当前脚本不是 Windows 服务管理器，也不提供自动开机拉起。")
lead(doc, "备份。", "在停止任务写入后，成套备份 app/backend/db.sqlite3、app/backend/media、Gis/kqmsdata、kqgis_publish、原始 GIS 数据与授权文件。至少保留一份异机备份，并记录备份时间、工程根路径及 KQGIS 服务状态。SQLite 与 media 应来自同一时点。")
lead(doc, "更新。", "先备份，再替换 app 中的代码并运行一键脚本执行迁移和重新启动。不要用新版空数据库覆盖正式 db.sqlite3，也不要清理 media、原始资料或 Gis/kqmsdata。若地图工程或原始数据位置变化，应单独进行 KQGIS 图层连接修复与 WMS 验收。")
lead(doc, "回退。", "停止应用写入，恢复同一批次的数据库与 media；如涉及 GIS 变更，恢复对应 kqmsdata/地图工程与原始资料后重新启动，并复查服务元数据和页面出图。不要将不同时间点的数据库、照片和 KQGIS 配置混用。")

doc.add_heading("8 常见故障定位", 1)
table(doc, ["现象", "优先检查", "处理方向"], [
    ("启动窗口很快关闭", "app/.runtime 最新 err.log；py、node、npm", "从 PowerShell 运行 scripts/start.ps1 读取报错；补齐依赖或路径。"),
    ("5173/5174 不可访问", "Get-NetTCPConnection；防火墙；BindAddress", "确认进程 ready，局域网放通前端端口并核对手机与服务器网络。"),
    ("地图灰底或专题图空白", "9600、KQGIS 授权、原始数据路径、服务元数据", "确认 jilong_port_multitemporal 和 jilong_county_atlas 正常出图；必要时修复并重新发布。"),
    ("登录后 403 或 CSRF 错误", "PublicOrigins、协议、Cookie Secure、代理头", "确保浏览器来源精确匹配；HTTPS 与安全 Cookie 配置一致，改后重启。"),
    ("DisallowedHost 或 Vite 拒绝域名", "PublicOrigins 与实际 Host", "把真实域名或 IP 来源写入配置，不带页面路径。"),
    ("队伍或任务消失", "db.sqlite3 是否为原库；迁移记录", "从同一批次备份恢复数据库；不要运行空库覆盖。"),
    ("照片无法显示", "media 路径、文件权限、会话权限", "核对照片文件与数据库记录，并通过已登录账号检查 /api 图片接口。"),
    ("端口 9600 已占用", "监听 PID 与实际程序", "确认是否本项目 KQGIS；勿盲目结束未知进程。"),
], [1.55, 2.1, 3.4])

doc.add_heading("9 交付前核对清单", 1)
bullets(doc, [
    "工程根路径、原始数据、KQGIS 地图工程及授权已在目标服务器核对。",
    "db.sqlite3 和 media 已成套迁移并备份；管理员与组长账号可用。",
    "四项服务和两个关键地图元数据接口响应正常；口岸及县域专题图可出图。",
    "局域网或最终域名在目标设备上可访问；/api 与 /kqgis 代理链路正常。",
    "公网仅开放预定的 HTTPS 入口；8000/9600 已从外部实测不可直连。",
    "演示口令、安全环境变量、HTTPS、备份和故障联系人已落实；若未完成生产化改造，明确限定演示访问范围。",
])

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
