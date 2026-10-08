from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\CangQiongPj")
OUT_DIR = ROOT / "artifacts" / "作品介绍文档_20260929"
ASSET_DIR = OUT_DIR / "assets"
OUTPUT = OUT_DIR / "吉隆口岸泥石流监测救援部署系统作品介绍.docx"

SONGTI = "宋体"
HEITI = "黑体"
ACCENT = "294E63"
ACCENT_LIGHT = "EAF2F5"
PALE = "F6F8F9"
BORDER = "D9D9D9"
TEXT = "1F2933"
MUTED = "52606D"


def font_path(name):
    return str(Path(r"C:\Windows\Fonts") / name)


def pil_font(name, size):
    return ImageFont.truetype(font_path(name), size)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, name=SONGTI, size=10.5, bold=False, color=TEXT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), name)


def set_paragraph_body(paragraph, indent=True, after=4, before=0, line=1.28):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if indent:
        fmt.first_line_indent = Pt(21)


def add_body(doc, text, indent=True, bold_lead=None, after=4):
    p = doc.add_paragraph()
    set_paragraph_body(p, indent=indent, after=after)
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, HEITI, 10.5, True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullet(doc, text, level=0, after=2):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.55 + level * 0.5)
    p.paragraph_format.first_line_indent = Cm(-0.25)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_number(doc, number, title, detail):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.15)
    p.paragraph_format.first_line_indent = Cm(-0.15)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.22
    r = p.add_run(f"{number}  {title}  ")
    set_run_font(r, HEITI, 10.5, True, ACCENT)
    r = p.add_run(detail)
    set_run_font(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(7 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    set_run_font(r, HEITI, 16 if level == 1 else 12, True, "000000")
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r, SONGTI, 9, False, MUTED)
    return p


def set_picture_alt(inline_shape, title, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_table(doc, headers, rows, widths=None, font_size=10.5, header_fill=ACCENT):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for index, header in enumerate(headers):
        cell = hdr.cells[index]
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell, 95, 110, 95, 110)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.1
        r = p.add_run(str(header))
        set_run_font(r, HEITI, font_size, True, "FFFFFF")
    for ridx, row_data in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for cidx, value in enumerate(row_data):
            cell = row.cells[cidx]
            if ridx % 2 == 1:
                set_cell_shading(cell, PALE)
            set_cell_margins(cell, 85, 105, 85, 105)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.16
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if cidx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, SONGTI, font_size)
        if widths:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    if widths:
        for idx, width in enumerate(widths):
            hdr.cells[idx].width = Cm(width)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, SONGTI, 9, False, MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def draw_centered(draw, box, text, font, fill="#1F2933", spacing=8):
    x0, y0, x1, y1 = box
    lines = text.split("\n")
    heights = [draw.textbbox((0, 0), line, font=font)[3] for line in lines]
    total = sum(heights) + spacing * (len(lines) - 1)
    y = y0 + (y1 - y0 - total) / 2
    for line, height in zip(lines, heights):
        bbox = draw.textbbox((0, 0), line, font=font)
        width = bbox[2] - bbox[0]
        draw.text((x0 + (x1 - x0 - width) / 2, y), line, font=font, fill=fill)
        y += height + spacing


def rounded(draw, box, fill, outline="#D4DEE3", radius=22, width=3):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def arrow(draw, start, end, color="#527A8C", width=5):
    draw.line([start, end], fill=color, width=width)
    x, y = end
    draw.polygon([(x, y), (x - 14, y - 9), (x - 14, y + 9)], fill=color)


def build_architecture_diagram(path):
    canvas = Image.new("RGB", (1600, 610), "white")
    draw = ImageDraw.Draw(canvas)
    title_font = pil_font("simhei.ttf", 34)
    box_font = pil_font("simhei.ttf", 29)
    small_font = pil_font("simsun.ttc", 23)
    draw.text((60, 28), "系统总体架构", font=title_font, fill="#111827")

    columns = [
        (55, 110, 450, 545, "数据资源层", "#F5F8FA"),
        (605, 110, 995, 545, "服务与业务层", "#EEF5F7"),
        (1150, 110, 1545, 545, "应用展现层", "#F5F8FA"),
    ]
    for x0, y0, x1, y1, label, fill in columns:
        rounded(draw, (x0, y0, x1, y1), fill, outline="#B9C8CE", radius=26)
        draw.text((x0 + 24, y0 + 21), label, font=box_font, fill="#294E63")

    source_items = [
        (95, 195, 410, 275, "赛事原始 GIS 数据"),
        (95, 305, 410, 385, "SQLite 业务数据"),
        (95, 415, 410, 495, "现场上报图片"),
    ]
    service_items = [
        (645, 195, 955, 300, "KQGIS Server\nWMS 地图服务"),
        (645, 340, 955, 475, "Django 5\n认证 任务 数据接口"),
    ]
    app_items = [
        (1190, 185, 1505, 275, "Vue 3 管理员端"),
        (1190, 310, 1505, 400, "OpenLayers 智慧大屏"),
        (1190, 435, 1505, 525, "救援组长移动端"),
    ]
    for box in source_items:
        rounded(draw, box[:4], "#FFFFFF")
        draw_centered(draw, box[:4], box[4], small_font)
    for box in service_items:
        rounded(draw, box[:4], "#FFFFFF", outline="#9EB7C2")
        draw_centered(draw, box[:4], box[4], small_font)
    for box in app_items:
        rounded(draw, box[:4], "#FFFFFF")
        draw_centered(draw, box[:4], box[4], small_font)
    arrow(draw, (450, 285), (600, 285))
    arrow(draw, (450, 395), (600, 395))
    arrow(draw, (995, 285), (1145, 285))
    arrow(draw, (995, 395), (1145, 395))
    canvas.save(path)


def build_flow_diagram(path):
    canvas = Image.new("RGB", (1600, 430), "white")
    draw = ImageDraw.Draw(canvas)
    title_font = pil_font("simhei.ttf", 33)
    stage_font = pil_font("simhei.ttf", 25)
    small_font = pil_font("simsun.ttc", 20)
    draw.text((55, 25), "救援任务状态流转", font=title_font, fill="#111827")
    boxes = [
        (45, 135, 285, 325, "管理员落点", "红色水滴\n待分配"),
        (365, 135, 605, 325, "分配具体队伍", "生成唯一编号\n进度 0%"),
        (685, 135, 925, 325, "组长到场上报", "任务中\n进度 50%"),
        (1005, 135, 1245, 325, "组长完成上报", "任务完成\n进度 100%"),
        (1325, 135, 1565, 325, "大屏同步结果", "绿底类别图标\n图片可查"),
    ]
    fills = ["#FFF4F4", "#FFF4F4", "#FFF4F4", "#EDF8F0", "#EDF8F0"]
    outlines = ["#D4555E", "#D4555E", "#D4555E", "#4F9B68", "#4F9B68"]
    for box, fill, outline in zip(boxes, fills, outlines):
        rounded(draw, box[:4], fill, outline=outline, radius=24, width=4)
        x0, y0, x1, y1 = box[:4]
        draw_centered(draw, (x0, y0 + 18, x1, y0 + 88), box[4], stage_font, fill="#1F2933")
        draw.line([(x0 + 24, y0 + 96), (x1 - 24, y0 + 96)], fill="#D7DEE2", width=2)
        draw_centered(draw, (x0, y0 + 104, x1, y1 - 10), box[5], small_font, fill="#52606D")
    for index in range(4):
        start = (boxes[index][2] + 8, 230)
        end = (boxes[index + 1][0] - 8, 230)
        arrow(draw, start, end, color="#527A8C", width=5)
    canvas.save(path)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = SONGTI
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), SONGTI)
    normal.paragraph_format.line_spacing = 1.28
    normal.paragraph_format.space_after = Pt(4)

    title = styles["Title"]
    title.font.name = HEITI
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), HEITI)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for level, size in ((1, 16), (2, 12)):
        style = styles[f"Heading {level}"]
        style.font.name = HEITI
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), HEITI)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    r = header.add_run("吉隆口岸泥石流监测救援部署系统  作品介绍")
    set_run_font(r, SONGTI, 8.5, False, MUTED)
    add_page_number(section.footer.paragraphs[0])


def page_break(doc):
    doc.add_page_break()


def build_document():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    architecture = ASSET_DIR / "系统架构图.png"
    flow = ASSET_DIR / "任务流转图.png"
    mobile_crop = ASSET_DIR / "组长任务端_文档裁剪.png"
    build_architecture_diagram(architecture)
    build_flow_diagram(flow)
    with Image.open(ASSET_DIR / "组长任务端.png") as mobile_image:
        mobile_image.crop((0, 0, 415, 760)).save(mobile_crop)

    doc = Document()
    configure_document(doc)

    # Page 1: cover and overview
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("吉隆口岸泥石流监测救援部署系统作品介绍")
    set_run_font(r, HEITI, 22, True, "000000")
    title_p_pr = p._p.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(26)
    r = p.add_run("苍穹杯全国大学生空间信息技术大赛")
    set_run_font(r, SONGTI, 13, False, MUTED)

    add_heading(doc, "作品摘要", 1)
    add_body(
        doc,
        "本作品以吉隆口岸泥石流灾害为典型场景，建设集空间数据展示、灾情点位标记、救援队伍调度和现场进度上报于一体的 Web 系统。系统基于赛事提供的 KQGIS Desktop 与 KQGIS Server 进行二次开发，将吉隆口岸多时相遥感影像和吉隆县专题空间数据发布为标准地图服务，并在管理员大屏与救援组长端之间建立任务闭环。",
    )
    add_body(
        doc,
        "作品以“不派生赛事空间数据”为数据约束。大屏中的冰川、冰湖、冻土和历史灾害区指标直接读取原始 CSV 与 Shapefile 字段；地图直接引用已发布的原始影像和专题图层。管理员落点后，系统可生成唯一任务编号并分配到具体队伍；组长用手机号登录，上传现场图片并推进任务阶段，大屏自动同步进度和点位图标。",
    )

    add_table(
        doc,
        ["项目", "内容"],
        [
            ("服务对象", "管理员、五类救援队伍组长"),
            ("核心范围", "吉隆口岸多时相影像与吉隆县专题图"),
            ("技术体系", "Vue 3 + OpenLayers + Django 5 + SQLite + KQGIS Server"),
            ("业务结果", "形成“监测展示、地图落点、队伍分配、现场上报、大屏回显”闭环"),
        ],
        widths=[3.2, 12.3],
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("关键词  泥石流  KQGIS  空间信息  救援调度  智慧大屏")
    set_run_font(r, SONGTI, 10.5, False, MUTED)

    # Page 2: requirements
    page_break(doc)
    add_heading(doc, "1 需求分析", 1)
    add_heading(doc, "1 1 场景与问题", 2)
    add_body(
        doc,
        "吉隆口岸及周边山高谷深，地形、冰川冰湖、冻土和河流等要素在灾情判读与救援组织中需要同时参考。传统做法中，空间数据、队伍名录和现场上报往往分散在不同载体中，管理员难以在同一视图内掌握“哪里需要处置、哪支队伍可用、任务已进展到哪一阶段”。本作品因此将空间可视化与救援任务流程统一设计。",
    )

    add_heading(doc, "1 2 用户需求", 2)
    add_table(
        doc,
        ["用户", "主要任务", "系统响应"],
        [
            ("管理员", "统筹地图、灾情点位和救援队伍", "查看原始数据与多时相影像，落点、分配、删除任务，维护队伍资料"),
            ("救援组长", "接收本组任务并反馈现场进展", "以组长手机号登录，查看任务编号与阶段，上传图片、上报下一阶段"),
        ],
        widths=[2.4, 5.1, 8.0],
        font_size=10,
    )

    add_heading(doc, "1 3 核心业务需求", 2)
    add_number(doc, "01", "空间数据统一展示", "首屏展示吉隆口岸影像，左侧卡片与吉隆县专题地图实现联动切换。")
    add_number(doc, "02", "灾情点位快速创建", "管理员直接在已发布影像范围内单击落点，未分配点以红色水滴图标表示。")
    add_number(doc, "03", "队伍可用性与分配", "大屏顶部按五类队伍显示“可用数与总数”，点位浮窗内完成具体队伍分配。")
    add_number(doc, "04", "任务闭环上报", "系统自动生成唯一任务编号，组长端分阶段上报进度和现场图片。")
    add_number(doc, "05", "结果回显与追溯", "大屏根据任务类别切换矢量图标，完成后改为绿底，悬停点位可查看阶段、进度和最新图片。")

    add_heading(doc, "1 4 数据与非功能约束", 2)
    add_table(
        doc,
        ["约束", "设计要求"],
        [
            ("数据真实性", "数据面板只展示原始文件已有记录数和字段，不计算新的风险分级或空间派生数据。"),
            ("可追溯性", "每个卡片标注源文件；地图图层与 KQGIS 服务及图层编号一一对应。"),
            ("响应与同步", "管理员端和组长端每 10 秒刷新业务状态，也在关键操作后立即刷新。"),
            ("可部署性", "两个前端可监听局域网与公网地址；Django 和 KQGIS 保留在内部，由前端反向代理访问。"),
        ],
        widths=[3.0, 12.5],
        font_size=9.7,
    )

    # Page 3: overall design
    page_break(doc)
    add_heading(doc, "2 总体设计", 1)
    add_heading(doc, "2 1 技术架构", 2)
    add_body(
        doc,
        "系统采用前后端分离架构。KQGIS Server 负责发布口岸影像、县域专题图和基础地图；Django 负责登录会话、队伍、点位、任务阶段、图片索引及原始数据读取接口；Vue 3 构建管理员端与组长端，OpenLayers 负责 WMS 地图叠加、交互落点与点位矢量样式。",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    picture = p.add_run().add_picture(str(architecture), width=Cm(15.7))
    set_picture_alt(
        picture,
        "系统总体架构",
        "数据资源层通过 KQGIS Server 和 Django 服务层连接管理员端、智慧大屏与救援组长移动端。",
    )
    add_caption(doc, "图 1  系统总体架构")

    add_heading(doc, "2 2 功能设计", 2)
    add_table(
        doc,
        ["功能域", "主要功能", "业务输出"],
        [
            ("登录与账号", "管理员登录、修改密码；组长手机号登录", "建立管理端与队伍端独立会话"),
            ("空间数据大屏", "9 景口岸影像切换；4 类专题数据卡片及县域地图联动", "提供可切换、可追溯的监测视图"),
            ("地图调度", "落点、删除、队伍分配、点位状态与图片查看", "产生任务编号和完整处置记录"),
            ("队伍管理", "按五类队伍新增、查看、编辑、删除队伍资料", "大屏队伍总数与可用数同步更新"),
            ("组长任务端", "查看当前任务、上传现场图片、推进任务阶段", "将进度、完成状态和图片回传管理员端"),
        ],
        widths=[3.1, 7.2, 5.2],
        font_size=9.5,
    )

    add_heading(doc, "2 3 前端页面组织", 2)
    add_body(
        doc,
        "管理员端由登录页、智慧大屏和调度后台构成。大屏保持“左侧数据卡片、中央地图、顶部队伍”的信息结构；点击任意队伍类别进入调度后台。组长端保持单列表单，便于在手机上查看当前任务并完成上报。",
    )

    # Page 4: GIS and data
    page_break(doc)
    add_heading(doc, "3 空间数据与地图服务设计", 1)
    add_heading(doc, "3 1 数据使用原则", 2)
    add_body(
        doc,
        "系统将赛事原始资料作为只读数据源。KQGIS 工程直接引用 GeoTIFF 与 Shapefile，不裁剪、不重投影、不分类；后端直接读取 CSV 行与 DBF 记录数、字段值，不把原型图中的演示数字当作真实数据。这一原则保证了界面指标、专题图层与原始文件之间的一致性。",
    )

    add_heading(doc, "3 2 左侧数据卡片与地图联动", 2)
    add_table(
        doc,
        ["卡片", "原始展示值", "数据来源", "专题层"],
        [
            ("冰川资料", "170 条记录；名称字段样本", "冰川.shp", "图层 6"),
            ("冰湖资料", "1990 年 101 处；2020 年 145 处；表面积字段 7.273240010689932 km²", "冰湖统计.csv 与冰湖 2020.shp", "图层 5"),
            ("冻土资料", "7 条记录；全部记录的原始 Shape_Area", "吉隆县冻土.shp", "图层 4"),
            ("历史灾害区", "84 条 PolygonZ 记录；前 6 条原始 Area", "历史灾害区 Shapefile", "图层 3"),
        ],
        widths=[2.8, 5.2, 5.0, 2.2],
        font_size=9.2,
    )
    add_body(
        doc,
        "用户点击卡片时，中央视图从口岸影像切换为吉隆县专题地图；再次点击或点击“返回吉隆口岸”即恢复口岸影像。县域视图将原始 DEM 与当前专题层作为独立 WMS 图层叠加，避免地形图层遮盖专题要素。",
    )

    add_heading(doc, "3 3 KQGIS 服务组织", 2)
    add_table(
        doc,
        ["服务", "内容", "前端用途"],
        [
            ("jilong_port_\nmultitemporal", "7 景 Vantor、1 景 Landsat 9、1 景 Sentinel 2A，共 9 景", "口岸核心地图与多时相切换"),
            ("jilong_county_atlas", "历史灾害区、冻土、冰湖、冰川、县界、DEM", "四类卡片对应的县域专题图"),
            ("jilong_base", "河流水系与吉隆县界", "口岸和县域视图的可选基础要素"),
        ],
        widths=[4.2, 6.6, 4.7],
        font_size=9.3,
    )

    add_heading(doc, "3 4 坐标与交互边界", 2)
    add_body(
        doc,
        "前端以 EPSG 4326 展示地图坐标。新增点位时，后端校验经纬度是否位于已发布的任一口岸影像包围框内，防止在无影像支撑的位置创建任务。",
    )

    # Page 5: database
    page_break(doc)
    add_heading(doc, "4 数据库设计", 1)
    add_heading(doc, "4 1 设计目标", 2)
    add_body(
        doc,
        "空间原始数据由 KQGIS 和文件系统管理，SQLite 仅保存业务过程中产生的队伍、任务点位、阶段及图片索引。这种分工避免复制大体量栅格，也避免业务操作修改赛事原始资料。Django 自带的用户、会话和迁移表承担管理员认证与结构版本管理。",
    )

    add_heading(doc, "4 2 核心实体", 2)
    add_table(
        doc,
        ["实体", "作用", "主要字段", "关系"],
        [
            ("RescueTeam 救援队伍", "保存五类队伍及组长资料", "category、name、status、captain、captain_phone、leader_password、members、note", "队伍名唯一；非空组长手机号唯一"),
            ("MapMarker 任务点位", "保存管理员落点和任务状态", "longitude、latitude、group、task_number、status、progress、current_stage、created_by", "由管理员创建；分配后关联具体队伍名"),
            ("MarkerReport 任务上报", "记录每次图片或阶段上报", "marker_id、content、progress、completed、photo、created_at", "多条上报属于一个任务点位；删除点位时级联删除"),
            ("Django User 与 Session", "管理员登录、密码哈希与会话", "username、password、is_staff、session_key", "MapMarker.created_by 关联管理员"),
        ],
        widths=[3.3, 3.7, 5.2, 3.3],
        font_size=8.9,
    )

    add_heading(doc, "4 3 关键约束与编号策略", 2)
    add_bullet(doc, "救援队伍名称必须唯一；组长手机号规范化为数字后不允许被其他队伍重复使用。")
    add_bullet(doc, "组长密码以 Django 密码哈希保存，新增队伍的默认密码为 123，数据库不保存明文密码。")
    add_bullet(doc, "任务编号采用 JL 加创建日期加零填充点位 ID，例如 JL-20260925-000021；task_number 在数据库层唯一。")
    add_bullet(doc, "点位只允许从“待分配”分配一次；已进入执行中的点位不能重复分配。")
    add_bullet(doc, "进度不允许倒退；完成状态必须对应 100%；非本队任务的组长无权访问任务图片。")

    add_heading(doc, "4 4 实体关系", 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    for text, bold in [
        ("Django User", True), (" 1 ── N ", False), ("MapMarker", True),
        (" 1 ── N ", False), ("MarkerReport", True),
        ("     队伍名关联     ", False), ("RescueTeam", True),
    ]:
        r = p.add_run(text)
        set_run_font(r, HEITI if bold else SONGTI, 11, bold, ACCENT if bold else MUTED)
    add_body(
        doc,
        "点位与队伍之间保留稳定的队伍名关联；管理员修改队伍名时，后端在事务中同步更新已有点位的 group 字段，避免任务失联。",
    )

    # Page 6: key technologies
    page_break(doc)
    add_heading(doc, "5 关键技术", 1)
    add_heading(doc, "5 1 任务状态机与可视化编码", 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    picture = p.add_run().add_picture(str(flow), width=Cm(15.7))
    set_picture_alt(
        picture,
        "救援任务状态流转",
        "管理员落点后分配具体队伍，组长分阶段上报，任务进度由零提升到五十和一百，最终在大屏同步绿色完成图标。",
    )
    add_caption(doc, "图 2  任务状态、进度与地图图标联动")
    add_body(
        doc,
        "未分配点位保持水滴图标。分配后，图标切换为对应队伍类别的红底圆形矢量图标，完成后改为绿底。五类图标分别以救援工具、医疗十字、闪电、山体监测和安置房屋为语义，在高密度地图上仍可快速识别任务类型与状态。",
    )

    add_heading(doc, "5 2 空间服务与交互地图", 2)
    add_body(
        doc,
        "KQGIS Server 以 WMS 输出栅格与矢量图层，OpenLayers 在 EPSG 4326 视图中加载影像层、河流层和业务点位矢量层。影像切换只更新 WMS 的 LAYERS 参数；点位悬停浮窗使用地图像素坐标动态定位，将分配表单、进度和现场图片保持在点位附近。",
    )

    add_heading(doc, "5 3 原始数据轻量解析", 2)
    add_body(
        doc,
        "后端不依赖额外的空间分析平台生成统计。程序按 Shapefile 与 DBF 规范读取文件头记录数和已有字段，并读取 CSV 原始行。对 2020 年冰湖，运行时还校验 CSV 的 lake_count 是否与对应 Shapefile 记录数一致，不一致时拒绝输出该仪表数据。",
    )

    add_heading(doc, "5 4 认证、权限与文件安全", 2)
    add_body(
        doc,
        "管理员使用 Django 用户会话和 CSRF 校验，修改密码时校验原密码及密码验证规则。组长会话只绑定一支队伍，只能读取和上报分配给本队的任务。图片上传限制为 10 MB 以内的 JPEG、PNG 或 WebP，后端同时校验文件扩展名和文件头特征。",
    )

    add_heading(doc, "5 5 同步与部署", 2)
    add_body(
        doc,
        "管理员端和组长端使用相同 Django API 与数据库，定时刷新与操作后立即刷新结合，使新增队伍、任务分配、阶段上报及图片在两端保持一致。开发环境提供一键启动脚本，同时启动 KQGIS、Django 和两个 Vue 前端；正式上线时可以反向代理和 HTTPS 统一对外提供 80 或 443 端口。",
    )

    # Page 7: interface evidence
    page_break(doc)
    add_heading(doc, "6 作品界面与操作效果", 1)
    add_heading(doc, "6 1 管理员智慧大屏", 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    picture = p.add_run().add_picture(str(ASSET_DIR / "智慧大屏.png"), width=Cm(15.7))
    set_picture_alt(
        picture,
        "管理员智慧大屏",
        "页面左侧是四类真实数据卡片，顶部是五类救援队伍可用数，中央是带有任务点位的吉隆口岸遥感影像。",
    )
    add_caption(doc, "图 3  智慧大屏实际运行界面")
    add_body(
        doc,
        "图 3 显示系统实际运行状态。左侧卡片显示原始记录数、字段图表与来源；顶部按五类队伍显示可用数和总数；中央为吉隆口岸遥感影像，可从 9 景已发布影像中切换。地图上的水滴、红底队伍图标和绿底图标分别表示待分配、执行中和已完成。",
    )

    add_heading(doc, "6 2 救援组长移动端", 2)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="FFFFFF", size="0")
    set_repeat_table_header(table.rows[0])
    left, right = table.rows[0].cells
    left.width = Cm(5.0)
    right.width = Cm(10.2)
    left.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    right.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture = p.add_run().add_picture(str(mobile_crop), width=Cm(4.2))
    set_picture_alt(
        picture,
        "救援组长移动端",
        "手机端展示当前任务编号、组名、阶段、五十百分比进度和现场图片上传控件。",
    )
    p = right.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("当前任务一屏处理")
    set_run_font(r, HEITI, 11, True, ACCENT)
    mobile_points = [
        "组长登录后只看到本队任务，不需要在地图与表格之间切换。",
        "任务编号、组名、当前阶段和进度均为只读，避免现场人员误改调度信息。",
        "可仅上传现场图片，也可同时推进到下一阶段；两种操作都会回传管理员大屏。",
        "支持手机浏览器与局域网访问，界面为单列响应式布局。",
    ]
    for item in mobile_points:
        p = right.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.18
        r = p.add_run(item)
        set_run_font(r, SONGTI, 10)
    add_caption(doc, "图 4  组长移动端实际运行界面")

    # Page 8: highlights and conclusion
    page_break(doc)
    add_heading(doc, "7 作品亮点", 1)
    add_table(
        doc,
        ["亮点", "实现方式", "评审价值"],
        [
            ("真实数据与界面双向联动", "四类卡片只读取原始文件记录和字段，每张卡片可切换到对应专题地图", "数据可追溯，避免演示指标与地图内容脱节"),
            ("口岸与县域两种尺度", "口岸视图支持 9 景多时相影像，县域视图支持 DEM 与专题层叠加", "兼顾救援落点所需的局部细节与背景要素"),
            ("任务闭环而非静态大屏", "落点、分配、编号、上报、进度、图片与完成状态在同一业务模型中流转", "可直接演示从发现问题到处置完成的全过程"),
            ("点位附近快速调度", "分配表单、删除操作、任务进度与图片直接显示在点位悬浮窗", "降低管理员在地图和侧栏之间的视线切换"),
            ("队伍数量与可用性动态计算", "顶部分子为当前上线且未被执行中任务占用的队伍数，分母为该类别队伍总数", "新增、删除队伍或分配任务后无需手工修改大屏数字"),
            ("应用层与 GIS 层解耦", "业务状态保存在 Django 数据库，空间图层保持在 KQGIS 服务", "便于后续替换业务数据库、增加新图层或接入原生手机端"),
        ],
        widths=[3.1, 7.0, 5.4],
        font_size=8.9,
    )

    add_heading(doc, "8 完成度与应用价值", 1)
    add_body(
        doc,
        "当前作品已实现管理员登录与改密码、智慧大屏、KQGIS 地图切换、地图点位管理、救援队伍维护、任务分配、组长端图片与阶段上报、两端状态同步，并提供一键启动、前端构建检查和 Django 自动化测试。系统可用于比赛演示与局域网联调，上线时将端口、域名、可信来源与 HTTPS 按服务器环境配置即可对外访问。",
    )
    add_body(
        doc,
        "本作品的设计重点是让空间数据直接服务于救援组织：原始图层提供场景依据，任务点位承载调度指令，组长上报形成现场反馈，大屏再将结果回到同一张地图上。因此，系统既能展示吉隆口岸泥石流相关数据，也能支撑一次完整的救援调度演练。",
    )

    props = doc.core_properties
    props.title = "吉隆口岸泥石流监测救援部署系统作品介绍"
    props.subject = "苍穹杯全国大学生空间信息技术大赛评审材料"
    props.author = "项目团队"
    props.keywords = "KQGIS,泥石流,空间信息,救援调度,智慧大屏"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
